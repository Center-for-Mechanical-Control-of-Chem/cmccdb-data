"""Actual XLSX authoring and converter checks; fixtures contain no scientific ground truth."""

import copy
import json
import os
import subprocess
import tempfile
import unittest
import zipfile
from pathlib import Path

from cmccdb_extraction.export import spreadsheet_rows, header_rows, flatten, excel_backend_limits
from cmccdb_extraction.schema import canonical, dataset_from_json
from cmccdb_schema.dataset_constructor import DatasetConstructor
from cmccdb_schema.proto import reaction_pb2

NODE = os.environ.get("CMCCDB_XLSX_NODE")
TEST_ROOT = os.environ.get("CMCCDB_TEST_WORK_ROOT", str(Path.cwd()))


class BackendBudgetTests(unittest.TestCase):
    def test_large_export_budget_is_explicit_configurable_and_bounded(self):
        self.assertEqual(excel_backend_limits({}),dict(CMCCDB_XLSX_HEAP_MIB=8192,CMCCDB_XLSX_TIMEOUT_SECONDS=600))
        self.assertEqual(excel_backend_limits(dict(CMCCDB_XLSX_HEAP_MIB='2048',CMCCDB_XLSX_TIMEOUT_SECONDS='90')),
                         dict(CMCCDB_XLSX_HEAP_MIB=2048,CMCCDB_XLSX_TIMEOUT_SECONDS=90))
        for name,value in [('CMCCDB_XLSX_HEAP_MIB','16384'),('CMCCDB_XLSX_HEAP_MIB','511'),
                           ('CMCCDB_XLSX_HEAP_MIB','unbounded'),('CMCCDB_XLSX_TIMEOUT_SECONDS','0'),
                           ('CMCCDB_XLSX_TIMEOUT_SECONDS','1801')]:
            with self.subTest(name=name,value=value),self.assertRaisesRegex(ValueError,name):excel_backend_limits({name:value})


def fixture_reaction(index, extra=False):
    value = dict(reaction_id="cmcc-" + f"{index:032x}",
        provenance=dict(is_mined=True, record_created=dict(time=dict(value="2026-10-07T12:00:00.000001+0000"),
            person=dict(name="Format fixture", email="fixture@example.invalid"))),
        inputs={"123": dict(components=[dict(identifiers=[dict(type="NAME", value="00123")],
            amount=dict(mass=dict(value=0 if index % 2 else 0.000003, units="GRAM")),
            reaction_role="REACTANT", is_limiting=False)])},
        conditions=dict(mechanochemistry=dict(type="BALL_MILL", frequency=dict(value=30, units="HERTZ"),
            geometry=["NO", "00123", "kneading"], number_of_balls=0)),
        outcomes=[dict(products=[dict(identifiers=[dict(type="NAME", value="FALSE")], is_desired_product=False,
            reaction_role="PRODUCT", measurements=[dict(type="YIELD", percentage=dict(value=0))])])],
        notes=dict(procedure_details="NO"))
    if extra:
        value["inputs"]["456"] = dict(components=[dict(identifiers=[dict(type="NAME", value="extra")],
            amount=dict(mass=dict(value=2, units="GRAM")), reaction_role="ADDITIVE")])
        value["outcomes"][0]["products"][0]["features"] = {
            "numeric-text": dict(string_value="123", format="text"),
            "bytes": dict(bytes_value="AP8=", format="binary")}
    return value


@unittest.skipUnless(NODE, "Set CMCCDB_XLSX_NODE and NODE_PATH for actual XLSX integration")
class ExportTests(unittest.TestCase):
    def write(self, directory, data):
        rows, cells, headers = spreadsheet_rows(data)
        plan = dict(sheets=[dict(name="ReactionData", rows=rows, header_rows=headers, widths=[12]*len(rows[0]))])
        path = Path(directory)
        (path / "plan.json").write_text(canonical(plan))
        backend = Path(__file__).parents[1] / "src/cmccdb_extraction/write_workbook.mjs"
        proc = subprocess.run([NODE, str(backend), str(path / "plan.json"), str(path / "test.xlsx")],
                              capture_output=True, text=True, timeout=180)
        self.assertEqual(proc.returncode, 0, proc.stderr)
        original = dataset_from_json(data)
        actual = DatasetConstructor.enumerate_spreadsheet(str(path / "test.xlsx"), name=original.name,
            id=original.dataset_id.removeprefix("cmcc_dataset-"))
        self.assertEqual(len(actual.reactions), len(original.reactions))
        self.assertEqual({r.reaction_id: r for r in actual.reactions}, {r.reaction_id: r for r in original.reactions})
        return rows, cells

    def test_large_sheet_mixed_headers_zero_false_numeric_text_bytes_and_lists(self):
        # No fabricated chemistry is published: these are format edge cases.
        dataset = dict(dataset_id="cmcc_dataset-" + "a"*32, name="Format fixture",
                       reactions=[fixture_reaction(i, extra=i%3==0) for i in range(1, 121)])
        with tempfile.TemporaryDirectory(prefix="xlsx-tests-", dir=TEST_ROOT) as directory:
            rows, cells = self.write(directory, dataset)
            self.assertEqual(sum(row[0] == "REACTION" for row in rows if row), 2)
            self.assertGreater(len(cells), 3000)

    def test_small_sheet(self):
        dataset = dict(dataset_id="cmcc_dataset-" + "b"*32, name="Small fixture",
                       reactions=[fixture_reaction(1)])
        with tempfile.TemporaryDirectory(prefix="xlsx-tests-", dir=TEST_ROOT) as directory:
            self.write(directory, dataset)

    def test_chunked_values_preserve_literal_types_and_sparse_columns(self):
        from openpyxl import load_workbook
        dataset=dict(dataset_id='cmcc_dataset-'+'c'*32,name='Chunked format fixture',
                     reactions=[fixture_reaction(i) for i in range(1,301)])
        rows,_,headers=spreadsheet_rows(dataset)
        evidence=[['#!','Fixture source text','Zero','False','Empty slot','Date-like source label']]
        evidence += [['#!','00123' if n%2 else '=literal',0,False,'','2026-10-07T12:00:00'] for n in range(700)]
        plan=dict(sheets=[dict(name='ReactionData',rows=rows,header_rows=headers,widths=[12]*max(map(len,rows))),
                          dict(name='Evidence',rows=evidence,widths=[4,20,12,12,12,30],header_rows=[1])])
        with tempfile.TemporaryDirectory(prefix='xlsx-chunk-tests-',dir=TEST_ROOT) as directory:
            path=Path(directory);(path/'plan.json').write_text(canonical(plan))
            backend=Path(__file__).parents[1]/'src/cmccdb_extraction/write_workbook.mjs'
            proc=subprocess.run([NODE,str(backend),str(path/'plan.json'),str(path/'test.xlsx')],capture_output=True,text=True,timeout=180)
            self.assertEqual(proc.returncode,0,proc.stderr)
            actual=DatasetConstructor.enumerate_spreadsheet(str(path/'test.xlsx'),name=dataset['name'],id='c'*32)
            original=dataset_from_json(dataset)
            self.assertEqual({r.reaction_id:r for r in actual.reactions},{r.reaction_id:r for r in original.reactions})
            workbook=load_workbook(path/'test.xlsx',read_only=True,data_only=False)
            try:
                sheet=workbook['Evidence'];self.assertEqual(sum(1 for _ in sheet.iter_rows()),701)
                for row in [2,256,257,513,701]:
                    self.assertEqual(sheet.cell(row,2).data_type,'s')
                    self.assertEqual(sheet.cell(row,3).value,0)
                    self.assertIs(sheet.cell(row,4).value,False)
                    self.assertEqual(sheet.cell(row,6).value,'2026-10-07T12:00:00')
                profile=json.loads((path/'test.xlsx.backend.json').read_text());self.assertEqual(profile['stage'],'complete')
                self.assertTrue(profile['literalOnly'])
                self.assertTrue(profile['recalculated'])
            finally:workbook.close()


if __name__ == "__main__":
    unittest.main()
